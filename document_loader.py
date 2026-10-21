"""
kaynak/ klasöründeki PDF, Word, TXT ve JSON dosyalarını okur.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

import config


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".json", ".md"}


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # Tablolardaki metni de al
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)



def _read_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, ensure_ascii=False, indent=2)


def load_file(path: Path) -> str | None:
    ext = path.suffix.lower()
    try:
        if ext in {".txt", ".md"}:
            return _read_txt(path)
        if ext == ".pdf":
            return _read_pdf(path)
        if ext == ".docx":
            return _read_docx(path)
        if ext == ".json":
            return _read_json(path)
    except Exception as exc:  # noqa: BLE001 — dosya bazlı hata, diğerleri devam etsin
        print(f"[uyarı] {path.name} okunamadı: {exc}")
    return None


def kaynak_fingerprint(kaynak_dir: str | None = None) -> str:
    """
    Klasördeki desteklenen dosyaların ad+boyut+mtime özeti.
    Değişince RAG yeniden yüklenmeli.
    """
    folder = Path(kaynak_dir or config.KAYNAK_DIR)
    if not folder.exists():
        return ""
    parts: list[str] = []
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            st = path.stat()
            parts.append(f"{path.name}:{st.st_size}:{st.st_mtime_ns}")
        except OSError:
            continue
    return "|".join(parts)


def load_all_documents(kaynak_dir: str | None = None) -> list[dict]:
    """
    Her belge için {"source": dosya_adı, "text": içerik} döner.
    """
    folder = Path(kaynak_dir or config.KAYNAK_DIR)
    folder.mkdir(parents=True, exist_ok=True)

    documents: list[dict] = []
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        text = load_file(path)
        if text and text.strip():
            documents.append({"source": path.name, "text": text.strip()})
        else:
            print(f"[uyarı] {path.name}: metin çıkarılamadı (tarama PDF olabilir).")
    return documents
